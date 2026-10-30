"""生成「作品介绍文档」：对应报名表单 6 的七项要求，含界面截图。"""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
GRAY = RGBColor(0x60, 0x64, 0x6B)

document = Document()
section = document.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.2)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)


def set_font(style, name, size, bold=False, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), name)


set_font(document.styles["Normal"], BODY_FONT, 10.5)
document.styles["Normal"].paragraph_format.line_spacing = 1.5
document.styles["Normal"].paragraph_format.space_after = Pt(4)
for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
    set_font(document.styles[name], HEAD_FONT, size, bold=True, color=RGBColor(0x18, 0x18, 0x1B))
    document.styles[name].paragraph_format.space_before = Pt(12)
    document.styles[name].paragraph_format.space_after = Pt(5)


def body(text, bold=False, indent=True, size=10.5):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74) if indent else Cm(0)
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    return paragraph


def bullet(text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text).font.size = Pt(10.5)
    return paragraph


def table(headers, rows, widths):
    result = document.add_table(rows=1, cols=len(headers))
    result.style = "Table Grid"
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.font.bold = True
        run.font.size = Pt(9.5)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), "F4F4F5")
        cell._tc.get_or_add_tcPr().append(shading)
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            run.font.size = Pt(9.5)
    result.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    result._tbl.tblPr.append(layout)
    for index, width in enumerate(widths):
        result.columns[index].width = Cm(width)
        for row in result.rows:
            row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return result


def figure(path, caption, width=15.0):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(path, width=Cm(width))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    run = caption_paragraph.add_run(caption)
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def footer_page_number():
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind in ("begin", "instr", "separate", "end"):
        if kind == "instr":
            element = OxmlElement("w:instrText")
            element.text = " PAGE "
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


footer_page_number()

# ---------- 封面 ----------
title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(90)
run = title.add_run("简历工坊")
run.font.name = HEAD_FONT
run.font.size = Pt(34)
run.font.bold = True
run._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("作品介绍文档")
run.font.size = Pt(16)
run.font.color.rgb = ACCENT

meta = document.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run("传智杯 · AI Agent 智能助手 · 2026-10-30")
run.font.size = Pt(11)
run.font.color.rgb = GRAY

document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 一 ----------
document.add_heading("一、项目详情及使用方式", level=1)
body(
    "简历工坊是一个「本地优先」的求职工作台：把「看懂岗位要求 → 判断匹配 → 定向修改简历 → "
    "生成求职信 → 记录投递」串成一条可解释的链路。数据全部保存在访问者自己的浏览器里，"
    "服务端只承担 AI 代理，不存简历、不收集联系方式。"
)
body("使用流程（约 3 分钟）：", bold=True, indent=False)
body("① 打开首页 → 顶栏点「演示场景」，载入示例简历与 3 个目标岗位；")
body("② 在「待投清单」里看按匹配度排序的岗位，点「准备」；")
body("③ 弹窗里点「启动求职 Agent」：模型先规划步骤，前两步本地执行（解析岗位、诊断差距），"
     "第三步流式改写简介并做事实核对，第四步生成求职信；")
body("④ 点「采用并写入简介」，刷新后仍保留；不满意可一键撤销；")
body("⑤ 投完点「标记为已投」，在投递看板里跟进。")
body(
    "运行方式：本地 `npm install && npm run dev`（默认 http://localhost:3000），"
    "或使用在线演示地址（部署在 Vercel，配置一个 DEEPSEEK_API_KEY 即可）。"
    "浏览器扩展可选：从 `extension/` 加载已解压扩展，可在招聘网站抓取真实 JD。"
)
body(
    "适用人群：应届生与 1–5 年经验的转岗者——他们有明确方向，但缺简历优化经验，"
    "也说不清自己与岗位差在哪。"
)

# ---------- 二 ----------
document.add_heading("二、项目亮点", level=1)
body("1. 可解释的匹配，不是黑盒百分比", bold=True, indent=False)
body(
    "匹配分拆成「核心要求 55 + 加分技能 20 + 简历完整度 25」并逐项展示；"
    "核心要求按「同族技能任选其一」分组（React 与 Vue 不会同时被要求），"
    "缺口提示是「浏览器与性能优化」这种可执行的要求，而不是硬塞一个技能标签。"
)
body("2. AI 是流程主导者，而不是右上角一个按钮", bold=True, indent=False)
body(
    "求职 Agent 先由模型规划步骤（工具白名单校验），再逐步执行；本地步骤跑完会把观察回灌给模型重新规划，"
    "用户也可以补充要求再规划。每一步都有状态、耗时与工具名，形成可查的工具调用轨迹。"
)
body("3. 克制的 AI：能核对、可撤销、不替你编", bold=True, indent=False)
body(
    "生成后本地做事实核对（数字 / 技能 / 名称是否超出素材），结果明确展示但不阻断；"
    "必须点「采用」才写入简历，且写入进入撤销栈。没有 Key、断网或模型输出不合规时自动降级到本地计划。"
)
body("4. 真实岗位数据，来自浏览器扩展", bold=True, indent=False)
body(
    "Chrome MV3 扩展从 BOSS直聘 / 智联 / 猎聘等站点抓取当前岗位页的 JD，经本地消息桥进入待投清单。"
    "抓取采用 JSON-LD → 站点选择器 → 标题解析 → 正文兜底四层策略，并提供「诊断选择器」定位失效规则。"
    "只读当前页、不抓列表、不自动提交，规避平台条款与封号风险。"
)
body("5. 隐私优先：数据不出浏览器", bold=True, indent=False)
body(
    "简历与投递记录存在 localStorage，AI 只发送当前板块文本、岗位关键词与参考 JD，"
    "不发送姓名、电话、邮箱与整份简历；API Key 只存在于服务端环境变量。"
)
body("6. 交付完整：能打印、能过 ATS、能离线", bold=True, indent=False)
body(
    "8 套模板、字体搭配、中英双语、手动分页与页码；双栏模板打印时改为「浮动侧栏 + 固定栏宽正文」，"
    "解决分栏无法跨页的问题；ATS 导出为纯文本并按 JD 检查关键词覆盖；PWA 外壳可离线打开。"
)

# ---------- 三 ----------
document.add_heading("三、AI 技术落地使用情况", level=1)
body("AI 在产品中承担三类任务，边界清晰：", indent=False)
table(
    ["环节", "模型做什么", "不做什么"],
    [
        ["单段写作", "按岗位重写简介 / 经历 / 项目描述", "不新增经历、不编造数字"],
        ["整份生成", "把一段自述整理成结构化简历", "不生成姓名、联系方式"],
        ["求职规划", "决定接下来做哪几步、给出一句理由", "不直接改简历、不调用未授权的工具"],
    ],
    [3.0, 7.2, 6.0],
)
body("落地约束（工程上如何保证可用）：", bold=True, indent=False)
bullet("输入裁剪：JD 2000 字、技能 20 个、观察 8 条封顶，避免无意义的长上下文与费用。")
bullet("输出校验：规划结果必须是工具白名单内的 JSON，≤5 步、去重、超长截断；不合规直接丢弃。")
bullet("失败降级：模型异常时重试一次，仍失败则回退本地规则计划；没有 Key 时整条链路仍可演示。")
bullet("事实核对：每次生成后本地比对数字、技能、名称，列出待核对项。")
body("质量数据（可复现，仓库内 `npm run eval` 生成）：", bold=True, indent=False)
table(
    ["指标", "结果"],
    [
        ["岗位画像识别", "26 / 26"],
        ["核心要求分组收益（逐项 82% → 分组 100%）", "+18 个百分点"],
        ["5 段手写 JD 的画像识别 / 缺口提示", "5 / 5、5 / 5"],
        ["事实核对：注入 20 条素材外事实", "20 / 20 检出"],
        ["事实核对：5 条干净文本", "0 误报"],
        ["单元测试 / 浏览器端到端", "274 个用例、237 项检查"],
    ],
    [10.4, 5.8],
)

# ---------- 四 ----------
document.add_heading("四、关键技术", level=1)
table(
    ["技术", "在项目中的作用"],
    [
        ["Next.js 16 App Router + React 19", "页面静态渲染、AI 路由与页面同域（省去跨域与鉴权）"],
        ["TypeScript strict", "简历数据模型、匹配规则、Agent 步骤共享类型"],
        ["SSE 流式输出", "AI 文案逐字出现，长文本感知延迟显著下降"],
        ["DeepSeek Chat + 结构化输出", "短文本改写与规划，支持 JSON 模式与流式"],
        ["提示词工程", "角色约束、素材边界、语言标记 ===EN===、长度限制"],
        ["本地存储与多标签页一致性", "localStorage + storage 事件，未编辑页跟随、编辑页提示冲突"],
        ["CSS @page 打印", "分页控制、页码、双栏浮动降级"],
        ["Chrome MV3", "抓取真实 JD、填充表单、选择器诊断"],
        ["Vitest + Puppeteer", "单元 / 组件测试 + 生产构建上的端到端回归"],
    ],
    [5.6, 10.6],
)

document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 五 ----------
document.add_heading("五、项目图片", level=1)
body("以下截图均来自真实运行环境（演示场景 + 真实模型调用）。", indent=False)
figure("D:/web/docs/img/editor.png", "图 1  编辑页：左侧表单、右侧实时预览，可直接打印为 PDF")
figure("D:/web/docs/img/match.png", "图 2  职位推荐：匹配分、核心要求命中与可点击的缺口")
figure("D:/web/docs/img/agent.png", "图 3  求职 Agent：四步计划、流式输出、事实核对与工具轨迹")
figure("D:/web/docs/img/queue.png", "图 4  待投清单：批量导入的岗位按匹配度排序，右侧可直接「准备」")

document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 六 ----------
document.add_heading("六、发表与获奖情况", level=1)
body(
    "本项目为本次参赛首次提交，此前未在任何平台发表、未参加其他赛事、未获得奖项。"
    "项目代码为独立开发，未使用任何开源项目模板或脚手架二次包装；"
    "所用依赖均为公开框架与库（Next.js、React、Tailwind CSS、Vitest、Puppeteer 等），"
    "在 `package.json` 中如实列出，前端无任何付费组件或素材。"
)

# ---------- 七 ----------
document.add_heading("七、原创性与参考说明", level=1)
body(
    "本项目不是在其他作品基础上的二次开发，代码与设计均为原创。开发过程中参考的公开资料限于："
    "Next.js 官方文档（App Router、Route Handlers、打印相关 CSS）、DeepSeek 官方 API 文档（流式与 JSON 模式）、"
    "schema.org 的 JobPosting 规范（用于扩展解析招聘页的结构化数据）。"
    "简历模板与视觉设计为本项目自行实现，未使用第三方模板。"
)
body(
    "需要说明的是：产品定位与主流简历工具不同——它们聚焦「排版」，本项目聚焦「投什么、怎么投、投了没」，"
    "并通过浏览器扩展接入真实岗位数据、通过可解释的匹配与 Agent 工作流把 AI 放在流程主导位置。"
)

document.save("D:/web/提交材料/作品介绍文档.docx")
print("intro docx saved")
